from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "presentation"
OUT_FILE = OUT_DIR / "lectora-hackathon-defense.pptx"
SHOTS = ROOT / "docs" / "screenshots"

W, H = Inches(13.333), Inches(7.5)

BG = RGBColor(8, 11, 18)
SURFACE = RGBColor(18, 24, 38)
SURFACE_2 = RGBColor(25, 32, 50)
WHITE = RGBColor(248, 250, 252)
MUTED = RGBColor(148, 163, 184)
PURPLE = RGBColor(124, 58, 237)
PURPLE_2 = RGBColor(167, 139, 250)
CYAN = RGBColor(34, 211, 238)
GREEN = RGBColor(52, 211, 153)
RED = RGBColor(251, 113, 133)


def add_text(slide, text, x, y, w, h, size=22, color=WHITE, bold=False,
             font="Arial", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rect(slide, x, y, w, h, fill=SURFACE, radius=True, line=None, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line if line else fill
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def add_line(slide, x1, y1, x2, y2, color=MUTED, width=1.5):
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_picture_crop(slide, path, x, y, w, h):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    return pic


def base_slide(prs, title, number=None, eyebrow=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if eyebrow:
        add_text(slide, eyebrow.upper(), 0.72, 0.34, 5.5, 0.28, 10, CYAN, True)
    add_text(slide, title, 0.72, 0.66, 11.1, 1.02, 28, WHITE, True)
    add_rect(slide, 0.72, 1.58, 0.68, 0.055, PURPLE, False)
    if number is not None:
        add_text(slide, f"{number:02d}", 12.05, 0.35, 0.5, 0.3, 10, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def add_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.text = text


def pill(slide, text, x, y, w, color=PURPLE, text_color=WHITE):
    add_rect(slide, x, y, w, 0.4, color, True)
    add_text(slide, text, x, y + 0.015, w, 0.34, 12, text_color, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def metric_card(slide, label, value, x, y, w, accent=PURPLE):
    add_rect(slide, x, y, w, 1.25, SURFACE, True, SURFACE_2)
    add_rect(slide, x, y, 0.07, 1.25, accent, False)
    add_text(slide, value, x + 0.28, y + 0.2, w - 0.45, 0.42, 22, WHITE, True)
    add_text(slide, label, x + 0.28, y + 0.7, w - 0.45, 0.32, 12, MUTED)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 — title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    add_rect(s, 8.55, 0.0, 4.783, 7.5, PURPLE, True, transparency=12)
    add_rect(s, 9.18, 0.28, 3.5, 6.8, SURFACE_2, True, transparency=2)
    add_picture_crop(s, SHOTS / "03-lecture-search.png", 9.35, 0.52, 3.16, 1.98)
    add_picture_crop(s, SHOTS / "06-player-subtitles.png", 9.35, 2.72, 3.16, 1.98)
    add_picture_crop(s, SHOTS / "05-account-groups.png", 9.35, 4.92, 3.16, 1.98)
    pill(s, "TEAMMENTALITY", 0.72, 0.55, 1.86, CYAN, BG)
    add_text(s, "Lectora превращает видео\nв знания за один\nрабочий процесс", 0.72, 1.35, 7.2, 2.45, 36, WHITE, True, line_spacing=0.94)
    add_text(s, "Поиск · конспект · перевод · субтитры · озвучка", 0.76, 4.05, 7.2, 0.44, 18, PURPLE_2, True)
    add_text(s, "Ковалёв Юрий Юрьевич\nОрлов Александр Дмитриевич\nТруфанова Валерия Витальевна", 0.76, 5.0, 5.8, 1.25, 16, WHITE, False, line_spacing=1.25)
    add_text(s, "Защита проекта · 5 минут", 0.76, 6.72, 4.0, 0.3, 11, MUTED)
    add_notes(s, "Здравствуйте! Мы — команда TeamMentality. Представляем Lectora: платформу, которая превращает обычное видео в структурированный, доступный для поиска и перевода источник знаний. За пять минут покажем проблему, решение, техническую основу и уже работающий MVP.")

    # 2 — problem
    s = base_slide(prs, "Видео хранит ответы, но заставляет искать их вручную", 2, "Проблема · 40 секунд")
    add_text(s, "Линейное видео", 0.8, 1.9, 2.8, 0.4, 20, WHITE, True)
    add_line(s, 3.3, 2.15, 4.5, 2.15, RED, 2.5)
    add_text(s, "нелинейный вопрос", 4.75, 1.9, 3.0, 0.4, 20, RED, True)
    for i, (num, head, body) in enumerate([
        ("01", "Перемотка", "Фрагмент ищут на глаз"),
        ("02", "Ручной конспект", "Заметки отнимают время"),
        ("03", "Несколько сервисов", "Перевод и хранение раздельны"),
    ]):
        x = 0.75 + i * 4.13
        add_rect(s, x, 3.0, 3.78, 2.42, SURFACE, True, SURFACE_2)
        add_text(s, num, x + 0.28, 3.27, 0.62, 0.35, 13, CYAN, True)
        add_text(s, head, x + 0.28, 3.82, 3.1, 0.4, 21, WHITE, True)
        add_text(s, body, x + 0.28, 4.42, 3.12, 0.72, 18, MUTED)
    add_text(s, "Итог: знания есть, доступа к ним нет", 0.76, 6.32, 7.0, 0.42, 21, PURPLE_2, True)
    add_notes(s, "Видео стало главным форматом обучения и обмена опытом, но оно остаётся линейным. Чтобы найти один ответ, человек перематывает запись, делает заметки вручную и переключается между несколькими сервисами. Чем длиннее библиотека, тем дороже становится повторное использование уже записанных знаний.")

    # 3 — solution
    s = base_slide(prs, "Lectora делает каждую запись доступной для поиска и обучения", 3, "Решение · 45 секунд")
    steps = [
        ("1", "Загрузить", "файл / ссылка"),
        ("2", "Распознать", "речь + таймкоды"),
        ("3", "Найти", "по смыслу"),
        ("4", "Применить", "перевод + конспект"),
    ]
    for i, (n, head, sub) in enumerate(steps):
        x = 0.78 + i * 3.03
        add_rect(s, x, 2.05, 2.55, 1.72, SURFACE, True, SURFACE_2)
        add_rect(s, x + 0.18, 2.22, 0.52, 0.52, PURPLE if i < 3 else CYAN, True)
        add_text(s, n, x + 0.18, 2.23, 0.52, 0.46, 15, BG if i == 3 else WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(s, head, x + 0.82, 2.2, 1.5, 0.42, 18, WHITE, True)
        add_text(s, sub, x + 0.82, 2.78, 1.52, 0.52, 18, MUTED)
        if i < 3:
            add_line(s, x + 2.56, 2.86, x + 2.96, 2.86, PURPLE_2, 2)
    add_picture_crop(s, SHOTS / "01-library.png", 0.78, 4.18, 5.3, 2.2)
    add_rect(s, 6.42, 4.18, 5.9, 2.2, SURFACE, True, SURFACE_2)
    add_text(s, "Одна платформа", 6.78, 4.53, 2.8, 0.4, 22, WHITE, True)
    add_text(s, "Библиотека и группы\nПрава доступа\nДва языка и озвучка", 6.8, 5.03, 4.8, 1.18, 18, MUTED, line_spacing=1.12)
    add_notes(s, "Lectora собирает весь путь в одном интерфейсе. Пользователь загружает файл или ссылку, получает транскрипт с таймкодами, ищет не по точному слову, а по смыслу, формирует конспект, перевод и озвучку. Результат хранится в личной библиотеке, группируется и при необходимости открывается другим по ссылке.")

    # 4 — tech
    s = base_slide(prs, "Локальный конвейер соединяет ASR, поиск и две LLM", 4, "Техническая реализация · 60 секунд")
    layers = [
        ("ПРОДУКТ", "React · Flask · PostgreSQL", CYAN),
        ("МЕДИА", "FFmpeg · Whisper / Vosk", GREEN),
        ("ПОИСК", "BGE-M3 · reranker", PURPLE_2),
        ("ГЕНЕРАЦИЯ", "DeepSeek API / Qwen", PURPLE),
    ]
    for i, (label, tech, accent) in enumerate(layers):
        y = 1.92 + i * 0.95
        add_rect(s, 0.78, y, 7.24, 0.76, SURFACE, True, SURFACE_2)
        add_rect(s, 0.78, y, 0.08, 0.76, accent, False)
        add_text(s, label, 1.08, y + 0.2, 1.78, 0.34, 15, accent, True)
        add_text(s, tech, 2.95, y + 0.17, 4.65, 0.4, 18, WHITE, True)
    add_rect(s, 8.38, 1.86, 3.94, 4.44, SURFACE, True, SURFACE_2)
    add_text(s, "Ускорение", 8.76, 2.2, 2.5, 0.38, 21, WHITE, True)
    metric_card(s, "транскрибация", "CUDA", 8.76, 2.86, 3.18, PURPLE)
    metric_card(s, "кодирование видео", "NVENC", 8.76, 4.25, 3.18, CYAN)
    add_text(s, "Выбор: локальная обработка или API", 8.78, 5.83, 3.08, 0.35, 13, MUTED, True)
    add_text(s, "Тяжёлые этапы готовы к выносу на воркеры", 0.8, 6.46, 10.8, 0.48, 19, PURPLE_2, True)
    add_notes(s, "Архитектуру проекта мы сохранили модульной. React-клиент общается с Flask API и PostgreSQL. FFmpeg готовит медиа, Whisper или Vosk распознаёт речь. BGE-M3 и reranker отвечают за семантический поиск. Для генерации можно выбрать DeepSeek API или локальную Qwen. CUDA ускоряет распознавание, NVENC — видео. Тяжёлые этапы изолированы, поэтому их можно вынести на отдельные воркеры без переделки интерфейса.")

    # 5 — MVP
    s = base_slide(prs, "MVP уже закрывает путь от загрузки до результата", 5, "Демонстрация · 60 секунд")
    add_picture_crop(s, SHOTS / "02-upload-translation.png", 0.76, 1.84, 5.9, 3.69)
    add_picture_crop(s, SHOTS / "06-player-subtitles.png", 6.76, 1.84, 5.82, 3.69)
    pill(s, "1 · Язык и перевод", 0.95, 5.68, 2.0, PURPLE)
    pill(s, "2 · Субтитры и озвучка", 6.95, 5.68, 2.35, CYAN, BG)
    add_text(s, "Далее: поиск ведёт к точному таймкоду, результат сохраняется в библиотеке", 0.8, 6.42, 11.4, 0.38, 18, WHITE, True)
    add_notes(s, "Демонстрация проходит в трёх действиях. Первое: загружаем видео, выбираем исходный язык, перевод и видимость. Второе: после обработки открываем плеер — переключаем русские, английские или двойные субтитры и озвучку в одном интерфейсе. Третье: вводим смысловой запрос и сразу переходим к нужному таймкоду. Видео остаётся в библиотеке, его можно добавить в группу или открыть по ссылке.")

    # 6 — advantages
    s = base_slide(prs, "Один сервис заменяет просмотр, перевод и ручной конспект", 6, "Ценность · 40 секунд")
    cards = [
        ("ПО СМЫСЛУ", "Поиск сразу ведёт к нужному моменту", CYAN),
        ("ДВА РЕЖИМА", "Локальная модель или облачный API", PURPLE_2),
        ("ДОСТУП", "Группы, права и просмотр по ссылке", GREEN),
        ("ДВА ЯЗЫКА", "Перевод, субтитры и озвучка RU ↔ EN", PURPLE),
    ]
    for i, (head, body, accent) in enumerate(cards):
        x = 0.78 + (i % 2) * 6.02
        y = 1.92 + (i // 2) * 2.18
        add_rect(s, x, y, 5.62, 1.78, SURFACE, True, SURFACE_2)
        add_rect(s, x + 0.28, y + 0.28, 0.12, 1.12, accent, True)
        add_text(s, head, x + 0.62, y + 0.28, 4.55, 0.32, 12, accent, True)
        add_text(s, body, x + 0.62, y + 0.76, 4.55, 0.72, 18, WHITE, True)
    add_text(s, "Практическая ценность: меньше повторной работы, больше повторного использования знаний", 0.8, 6.4, 11.4, 0.42, 18, PURPLE_2, True)
    add_notes(s, "Главное отличие Lectora — не отдельная AI-функция, а цельный сценарий. Семантический поиск ведёт прямо к моменту. Локальный и облачный режимы дают выбор между приватностью и качеством. Права доступа позволяют использовать платформу и внутри команды, и публично. А мультиязычность делает одну запись полезной для более широкой аудитории.")

    # 7 — close
    s = base_slide(prs, "Lectora готова к пилоту и масштабированию после хакатона", 7, "Итог · 20 секунд")
    add_text(s, "Следующий шаг", 0.8, 1.92, 3.0, 0.4, 22, CYAN, True)
    add_text(s, "Пилот на реальной библиотеке видео", 0.8, 2.45, 6.2, 0.62, 29, WHITE, True)
    for i, (n, txt) in enumerate([
        ("01", "масштабировать обработку воркерами"),
        ("02", "добавить аналитику обучения"),
        ("03", "расширить языки и модели"),
    ]):
        y = 3.48 + i * 0.72
        add_text(s, n, 0.82, y, 0.52, 0.3, 12, PURPLE_2, True)
        add_text(s, txt, 1.48, y - 0.04, 5.6, 0.48, 18, WHITE)
    add_rect(s, 8.0, 1.88, 4.26, 4.62, PURPLE, True)
    add_text(s, "Lectora", 8.48, 2.46, 3.2, 0.55, 30, WHITE, True)
    add_text(s, "Видео заканчивается.\nЗнания остаются.", 8.48, 3.36, 3.2, 1.2, 26, WHITE, True, line_spacing=1.0)
    add_text(s, "TeamMentality", 8.5, 5.66, 2.8, 0.35, 14, WHITE, True)
    add_text(s, "Спасибо! Готовы ответить на вопросы.", 0.8, 6.52, 6.4, 0.38, 18, MUTED)
    add_notes(s, "MVP готов к пилоту на реальной библиотеке видео. Дальше мы масштабируем обработку через воркеры, добавим аналитику обучения и расширим набор языков и моделей. Lectora: видео заканчивается — знания остаются. Спасибо, готовы ответить на вопросы.")

    prs.core_properties.title = "Lectora — защита проекта TeamMentality"
    prs.core_properties.subject = "Презентация на 5 минут для жюри хакатона"
    prs.core_properties.author = "TeamMentality"
    prs.core_properties.keywords = "Lectora, TeamMentality, AI, видео, семантический поиск"
    prs.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build()
